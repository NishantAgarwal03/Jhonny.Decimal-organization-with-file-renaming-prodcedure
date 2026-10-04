"""NLPEngine.extract_text: txt / pdf / docx / unsupported / damaged inputs. No model needed."""
import io
import os
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout

import docx

import fakes  # noqa: F401  (puts the repo root on sys.path)
from johnny.core.nlp_engine import NLPEngine

extract = lambda path: NLPEngine.extract_text(None, path)


def minimal_pdf(text):
    """A hand-built one-page PDF with a real text layer."""
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 50 Td ({text}) Tj ET"
    objs[3] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
    out, offsets = "%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


class ExtractText(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="johnny_test_")
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def path(self, name, data=b""):
        p = os.path.join(self.tmp, name)
        with open(p, "wb") as f:
            f.write(data)
        return p

    def test_txt_utf8_and_uppercase_extension(self):
        self.assertEqual(extract(self.path("a.TXT", "Aadhaar café".encode("utf-8"))), "Aadhaar café")

    def test_txt_invalid_bytes_do_not_crash(self):
        self.assertIn("Tax", extract(self.path("a.txt", b"Tax \xff\xfe return")))

    def test_empty_txt(self):
        self.assertEqual(extract(self.path("a.txt")), "")

    def test_pdf_text_layer(self):
        self.assertIn("Invoice", extract(self.path("a.pdf", minimal_pdf("Invoice Amazon 2023"))))

    def test_docx_paragraphs(self):
        p = os.path.join(self.tmp, "a.docx")
        d = docx.Document()
        d.add_paragraph("Salary slip")
        d.add_paragraph("March 2022")
        d.save(p)
        self.assertEqual(extract(p).split(), ["Salary", "slip", "March", "2022"])

    def test_unsupported_extension_returns_empty(self):
        self.assertEqual(extract(self.path("a.jpg", b"\xff\xd8\xff")), "")
        self.assertEqual(extract(self.path("noext", b"text")), "")

    def test_damaged_pdf_and_docx_return_empty_without_raising(self):
        with redirect_stdout(io.StringIO()):
            self.assertEqual(extract(self.path("bad.pdf", b"not a pdf")), "")
            self.assertEqual(extract(self.path("bad.docx", b"not a zip")), "")

    def test_extraction_never_modifies_the_file(self):
        p = self.path("a.txt", b"hello")
        mtime = os.path.getmtime(p)
        extract(p)
        with open(p, "rb") as f:
            self.assertEqual((f.read(), os.path.getmtime(p)), (b"hello", mtime))


if __name__ == "__main__":
    unittest.main()
